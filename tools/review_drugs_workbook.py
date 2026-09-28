from pathlib import Path
import re
from openpyxl import load_workbook

source = Path('public/drugs.xlsx')
target = Path('public/drugs-review.xlsx')
workbook = load_workbook(source)
sheet = workbook['Medicines List']
sheet.cell(1, 6, 'Review Status')
sheet.cell(1, 7, 'Review Notes')

safe_replacements = [
    (r'\bmeg\b', 'mcg'),
    (r'\bmp\b', 'mg'),
    (r'\bmceg\b', 'mcg'),
    (r'\bTodine\b', 'Iodine'),
    (r'\bTodized\b', 'Iodized'),
    (r'\bKrgocalciferol\b', 'Ergocalciferol'),
    (r'\bbydrochloride\b', 'hydrochloride'),
]

structural_rows = set(range(592, 596)) | set(range(636, 655)) | set(range(712, 719)) | {152, 154, 167, 523}
for row in range(2, sheet.max_row + 1):
    notes = []
    for column in range(1, 6):
        cell = sheet.cell(row, column)
        if isinstance(cell.value, str):
            original = cell.value
            for pattern, replacement in safe_replacements:
                cell.value = re.sub(pattern, replacement, cell.value, flags=re.IGNORECASE)
            if cell.value != original:
                notes.append('safe OCR unit/spelling correction')
    if row in structural_rows:
        sheet.cell(row, 6, 'MANUAL_REVIEW')
        notes.append('merged headings or medicine/formulation data requires manual reconstruction')
    else:
        sheet.cell(row, 6, 'CHECKED')
    sheet.cell(row, 7, '; '.join(dict.fromkeys(notes)))

workbook.save(target)
print(f'Created {target} with {sheet.max_row - 1} data rows.')
